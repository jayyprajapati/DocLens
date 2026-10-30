"""Build the PDF and DOCX variants of the eval corpus from the Markdown sources.

The Markdown files are the source of truth; this renders two of them to real
binary formats (with real tables) so the eval also exercises Brain's PDF/DOCX
extraction path. Needs `reportlab` and `python-docx` (see requirements-dev.txt).
"""
from __future__ import annotations

import re
from pathlib import Path

CORPUS = Path(__file__).parent / "corpus"


def _blocks(md: str):
    """Yield ('h1'|'h2'|'p'|'li'|'table', payload) from simple Markdown."""
    lines, i = md.splitlines(), 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("# "):
            yield "h1", ln[2:].strip()
        elif ln.startswith("## "):
            yield "h2", ln[3:].strip()
        elif ln.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(set(c) <= set("- :") for c in cells):
                    rows.append(cells)
                i += 1
            i -= 1
            yield "table", rows
        elif re.match(r"^(- |\d+\. )", ln):
            yield "li", re.sub(r"^(- |\d+\. )", "", ln).strip()
        elif ln.strip():
            yield "p", ln.strip()
        i += 1


def _plain(s: str) -> str:
    return re.sub(r"[*`]", "", s)


def build_docx(src: Path, dst: Path) -> None:
    import docx

    d = docx.Document()
    for kind, val in _blocks(src.read_text()):
        if kind == "h1":
            d.add_heading(_plain(val), level=1)
        elif kind == "h2":
            d.add_heading(_plain(val), level=2)
        elif kind == "li":
            d.add_paragraph(_plain(val), style="List Bullet")
        elif kind == "table":
            t = d.add_table(rows=len(val), cols=len(val[0]))
            t.style = "Table Grid"
            for r, row in enumerate(val):
                for c, cell in enumerate(row):
                    t.cell(r, c).text = _plain(cell)
        else:
            d.add_paragraph(_plain(val))
    d.save(dst)


def build_pdf(src: Path, dst: Path) -> None:
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    story = []
    for kind, val in _blocks(src.read_text()):
        if kind in ("h1", "h2"):
            story += [Paragraph(_plain(val), ss["Heading1" if kind == "h1" else "Heading2"]), Spacer(1, 6)]
        elif kind == "li":
            story.append(Paragraph("• " + _plain(val), ss["BodyText"]))
        elif kind == "table":
            t = Table([[_plain(c) for c in row] for row in val])
            t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.5, colors.black),
                                   ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey)]))
            story += [t, Spacer(1, 10)]
        else:
            story += [Paragraph(_plain(val), ss["BodyText"]), Spacer(1, 6)]
    SimpleDocTemplate(str(dst), pagesize=(612, 792)).build(story)


if __name__ == "__main__":
    build_docx(CORPUS / "q3_financial_report.md", CORPUS / "q3_financial_report.docx")
    build_pdf(CORPUS / "kestrel_postmortem.md", CORPUS / "kestrel_postmortem.pdf")
    print("built q3_financial_report.docx, kestrel_postmortem.pdf")
